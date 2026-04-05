#!/usr/bin/env python3
"""Parse Row IQ Dummy Erg Scores Excel → src/data/ergData.ts
Produces full per-session starboard/port athlete arrays + athlete summaries."""

import re
import json
import sys
from datetime import datetime, timedelta
from collections import defaultdict, Counter
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import load_workbook

EXCEL_PATH = Path(__file__).parent.parent.parent / "Downloads" / "Row IQ Dummy Erg Scores.xlsx"
OUTPUT_PATH = Path(__file__).parent.parent / "src" / "data" / "ergData.ts"

NAME_MAP = {
    "Pfautch": "Pfautsch", "Andreucetti": "Andreuccetti", "Distefano": "DiStefano",
    "Unlod": "Unold", "Joregensen": "Jorgensen", "Mackinney": "MacKinney",
    "Luludzija": "Luledzija", "Ried": "Reid", "Ferrrera": "Ferrera",
    "Koudraeiv": "Koudraiev", "Koudriaev": "Koudraiev", "Barrel": "Barrell",
    "van Doren": "Van Doren", "B.Bottomley": "Bottomley",
}

NOISE = {"Starboard", "Port", "AVG", "DNF", "Overall", "Weight Adj", "Weight Adjusted",
         "OVERALL AVG", "Second 12k", "", "—-", "—", "Erg", "Bike"}

SPLIT_RE = re.compile(r"(\d+\.?\d*)\s*\((\d+)")

CATEGORY = {
    "steady_state": [r"^30'", r"2x20", r"2x22", r"3x15", r"3x17", r"6x10"],
    "intervals": [r"4x4km", r"3x3km", r"2x6k", r"2x5km", r"3x5", r"10.+9.+8.+7",
                  r"30.+20.+10", r"5k.+4k.+3k", r"7.+30", r"4x4'"],
    "threshold": [r"^5km$"],
    "triathlon": [r"triath"],
}


def fix_name(name):
    if not name:
        return None
    name = str(name).strip()
    name = NAME_MAP.get(name, name)
    if name in NOISE or len(name) <= 1:
        return None
    if name.startswith(("Bike", "DNF", "Erg", "Second")):
        return None
    if re.match(r"^[\d\.\(\)\s:]+$", name):
        return None
    return name


def parse_split(raw):
    if raw is None:
        return None
    raw = str(raw).strip()
    if not raw or raw.upper().startswith("DNF"):
        return None
    m = SPLIT_RE.search(raw)
    if not m:
        return None
    s = float(m.group(1))
    spm = int(m.group(2))
    if s > 60:
        return None
    return {"split": round(s, 1), "spm": spm}


def classify(workout, sheet_name):
    text = (workout or "").strip()
    combined = text.lower() + " " + sheet_name.lower()
    for cat, patterns in CATEGORY.items():
        for p in patterns:
            if re.search(p, combined, re.IGNORECASE):
                return cat
    return "steady_state"


def find_sections(ws, header_row_num):
    """Find starboard/port name+avg column pairs, stopping at Bike/Overall."""
    sections = []
    cur = None
    for cell in ws[header_row_num]:
        v = str(cell.value).strip() if cell.value else ""
        vl = v.lower()
        if any(k in vl for k in ("bike", "overall", "weight")):
            break
        if vl in ("starboard", "port"):
            cur = {"label": v, "side": "starboard" if vl == "starboard" else "port",
                   "name_col": cell.column, "avg_col": None}
            sections.append(cur)
        elif cur and cur["avg_col"] is None and vl in ("avg", "overall avg", "avg split"):
            cur["avg_col"] = cell.column
    return [s for s in sections if s["avg_col"]]


def parse_standard_sheet(ws, sheet_name):
    """Parse a non-triathlon sheet."""
    date_val = ws.cell(1, 1).value
    workout = str(ws.cell(1, 2).value or "").strip()

    if isinstance(date_val, datetime):
        date_str = date_val.strftime("%Y-%m-%d")
    else:
        date_str = str(date_val)[:10]

    cat = classify(workout, sheet_name)

    # Find header row (usually 3, but 5km has it at row 2)
    header_row = 3
    r3a = ws.cell(3, 1).value
    if r3a:
        r3s = str(r3a).strip().lower()
        if r3s not in ("starboard", "port", ""):
            r2a = ws.cell(2, 1).value
            if r2a and str(r2a).strip().lower() in ("starboard", "port"):
                header_row = 2
            else:
                header_row = 3

    sections = find_sections(ws, header_row)
    if not sections:
        # Fallback for sheets without explicit Starboard/Port headers (like 5km)
        # Check if row 2 or 3 data starts with names
        sections = [
            {"side": "starboard", "name_col": 1, "avg_col": 2},
            {"side": "port", "name_col": 4, "avg_col": 5},
        ]
        header_row = 2  # data starts at row 3

    data_start = header_row + 1
    starboard = []
    port = []

    for sec in sections:
        for r in range(data_start, (ws.max_row or data_start) + 1):
            raw_name = ws.cell(r, sec["name_col"]).value
            name = fix_name(raw_name)
            if not name:
                if raw_name is None or str(raw_name).strip() == "":
                    break
                continue
            raw_val = ws.cell(r, sec["avg_col"]).value
            split_data = parse_split(raw_val)
            is_dnf = raw_val and str(raw_val).strip().upper().startswith("DNF")
            if split_data:
                entry = {"name": name, "split": split_data["split"], "spm": split_data["spm"],
                         "rank": 0, "dnf": False}
                if sec["side"] == "starboard":
                    starboard.append(entry)
                else:
                    port.append(entry)
            elif is_dnf:
                entry = {"name": name, "split": 999, "spm": 0, "rank": 0, "dnf": True}
                if sec["side"] == "starboard":
                    starboard.append(entry)
                else:
                    port.append(entry)

    starboard.sort(key=lambda x: x["split"])
    port.sort(key=lambda x: x["split"])
    for i, a in enumerate(starboard):
        a["rank"] = i + 1
    for i, a in enumerate(port):
        a["rank"] = i + 1

    all_valid = [a["split"] for a in starboard + port if not a["dnf"]]
    if not all_valid:
        return None

    s_valid = [a["split"] for a in starboard if not a["dnf"]]
    p_valid = [a["split"] for a in port if not a["dnf"]]

    return {
        "id": date_str,
        "date": date_str,
        "workoutName": workout,
        "workoutCategory": cat,
        "starboard": starboard,
        "port": port,
        "teamAvgSplit": round(sum(all_valid) / len(all_valid), 1),
        "bestSplit": round(min(all_valid), 1),
        "worstSplit": round(max(all_valid), 1),
        "bestAthlete": min((a for a in starboard + port if not a["dnf"]), key=lambda x: x["split"])["name"],
        "starboardAvg": round(sum(s_valid) / len(s_valid), 1) if s_valid else 0,
        "portAvg": round(sum(p_valid) / len(p_valid), 1) if p_valid else 0,
        "participantCount": len([a for a in starboard + port if not a["dnf"]]),
    }


def parse_triathlon_sheet(ws, sheet_name):
    """Parse triathlon sheets with timedelta values."""
    date_val = ws.cell(1, 1).value
    workout = str(ws.cell(1, 2).value or "").strip()
    date_str = date_val.strftime("%Y-%m-%d") if isinstance(date_val, datetime) else str(date_val)[:10]

    header = list(ws[3])
    s_name_col = s_time_col = p_name_col = p_time_col = None
    exclude = ws.max_column + 1

    for cell in header:
        v = str(cell.value).strip() if cell.value else ""
        vl = v.lower()
        if any(k in vl for k in ("bike", "overall", "erg")):
            exclude = min(exclude, cell.column)
            continue
        if vl == "starboard":
            s_name_col = cell.column
        elif vl == "port":
            p_name_col = cell.column
        elif vl.startswith("time"):
            if s_name_col and not s_time_col:
                s_time_col = cell.column
            elif p_name_col and not p_time_col and cell.column < exclude:
                p_time_col = cell.column

    starboard = []
    port = []

    for r in range(4, (ws.max_row or 4) + 1):
        if s_name_col and s_time_col:
            name = fix_name(ws.cell(r, s_name_col).value)
            tv = ws.cell(r, s_time_col).value
            if name and isinstance(tv, timedelta):
                ts = int(tv.total_seconds())
                m, s = divmod(ts, 60)
                starboard.append({"name": name, "split": ts, "spm": 0, "rank": 0,
                                  "dnf": False, "timeFormatted": f"{m}:{s:02d}"})

        if p_name_col and p_time_col and p_name_col < exclude:
            name = fix_name(ws.cell(r, p_name_col).value)
            tv = ws.cell(r, p_time_col).value
            if name and isinstance(tv, timedelta):
                ts = int(tv.total_seconds())
                m, s = divmod(ts, 60)
                port.append({"name": name, "split": ts, "spm": 0, "rank": 0,
                             "dnf": False, "timeFormatted": f"{m}:{s:02d}"})

    starboard.sort(key=lambda x: x["split"])
    port.sort(key=lambda x: x["split"])
    for i, a in enumerate(starboard):
        a["rank"] = i + 1
    for i, a in enumerate(port):
        a["rank"] = i + 1

    all_t = [a["split"] for a in starboard + port]
    if not all_t:
        return None

    return {
        "id": date_str,
        "date": date_str,
        "workoutName": workout,
        "workoutCategory": "triathlon",
        "starboard": starboard,
        "port": port,
        "teamAvgSplit": round(sum(all_t) / len(all_t), 1),
        "bestSplit": min(all_t),
        "worstSplit": max(all_t),
        "bestAthlete": min(starboard + port, key=lambda x: x["split"])["name"],
        "starboardAvg": round(sum(a["split"] for a in starboard) / len(starboard), 1) if starboard else 0,
        "portAvg": round(sum(a["split"] for a in port) / len(port), 1) if port else 0,
        "participantCount": len(starboard) + len(port),
    }


def parse_sheet(ws, sheet_name):
    workout = str(ws.cell(1, 2).value or "").strip()
    cat = classify(workout, sheet_name)
    if cat == "triathlon":
        return parse_triathlon_sheet(ws, sheet_name)
    return parse_standard_sheet(ws, sheet_name)


def build_athletes(sessions):
    """Build athlete summaries from sessions (excluding triathlon from split averages)."""
    data = defaultdict(lambda: {"sessions": [], "splits": [], "sides": []})

    for session in sessions:
        if session["workoutCategory"] == "triathlon":
            continue
        for side in ("starboard", "port"):
            for a in session[side]:
                if a["dnf"]:
                    continue
                name = a["name"]
                data[name]["sessions"].append({
                    "sessionId": session["id"],
                    "date": session["date"],
                    "workoutName": session["workoutName"],
                    "workoutCategory": session["workoutCategory"],
                    "side": side,
                    "split": a["split"],
                    "spm": a["spm"],
                    "rank": a["rank"],
                    "of": session["participantCount"],
                })
                data[name]["splits"].append(a["split"])
                data[name]["sides"].append(side[0].upper())

    athletes = []
    for name, d in data.items():
        slist = sorted(d["sessions"], key=lambda x: x["date"])
        splits = d["splits"]
        avg = round(sum(splits) / len(splits), 1)
        best = round(min(splits), 1)
        best_date = min(d["sessions"], key=lambda x: x["split"])["date"]
        side_c = Counter(d["sides"])
        primary = "S" if side_c.get("S", 0) >= side_c.get("P", 0) else "P"

        trend = "stable"
        td = 0.0
        if len(splits) >= 6:
            early = sum(splits[:3]) / 3
            late = sum(splits[-3:]) / 3
            td = round(late - early, 1)
            trend = "improving" if td < -0.5 else ("declining" if td > 0.5 else "stable")

        # Category breakdown
        cat_d = defaultdict(list)
        for s in slist:
            cat_d[s["workoutCategory"]].append(s["split"])
        cat_breakdown = {}
        for c in ("steady_state", "intervals", "threshold"):
            if c in cat_d:
                cs = cat_d[c]
                cat_breakdown[c] = {"sessions": len(cs), "avgSplit": round(sum(cs) / len(cs), 1)}
            else:
                cat_breakdown[c] = {"sessions": 0, "avgSplit": None}

        sparkline = [{"date": s["date"], "split": s["split"]} for s in slist]

        athletes.append({
            "id": name.lower().replace(" ", "-").replace("'", ""),
            "name": name,
            "group": primary,
            "rank": 0,
            "avgSplit": avg,
            "bestSplit": best,
            "bestSplitDate": best_date,
            "sessionCount": len(slist),
            "trend": trend,
            "trendDelta": td,
            "consistency": round(len(slist) / 26 * 100),
            "categoryBreakdown": cat_breakdown,
            "sparklineData": sparkline,
            "sessions": slist,
        })

    athletes.sort(key=lambda x: x["avgSplit"])
    for i, a in enumerate(athletes):
        a["rank"] = i + 1
    return athletes


def main():
    print(f"Reading: {EXCEL_PATH}")
    wb = load_workbook(str(EXCEL_PATH), data_only=True)

    sessions = []
    for sheet_name in wb.sheetnames[1:]:
        ws = wb[sheet_name]
        print(f"  {sheet_name}...", end=" ")
        session = parse_sheet(ws, sheet_name)
        if session:
            sessions.append(session)
            is_tri = session["workoutCategory"] == "triathlon"
            print(f"{'TRI' if is_tri else 'OK'} | {session['participantCount']} athletes"
                  + (f" | avg={session['teamAvgSplit']}" if not is_tri else ""))
        else:
            print("SKIP (no data)")

    sessions.sort(key=lambda x: x["date"])

    # Deduplicate IDs
    seen = set()
    for s in sessions:
        if s["id"] in seen:
            s["id"] += "b"
        seen.add(s["id"])

    athletes = build_athletes(sessions)

    # Validation
    print(f"\n=== Validation ===")
    print(f"Sessions: {len(sessions)}")
    print(f"Athletes: {len(athletes)}")

    non_tri = [s for s in sessions if s["workoutCategory"] != "triathlon"]
    max_avg = max(s["teamAvgSplit"] for s in non_tri)
    print(f"Max team avg (non-tri): {max_avg}")

    phelps = next((a for a in athletes if a["name"] == "Phelps"), None)
    if phelps:
        print(f"Phelps: #{phelps['rank']}, avg={phelps['avgSplit']}, best={phelps['bestSplit']}, sessions={phelps['sessionCount']}")

    # Check for bad splits
    bad = 0
    for s in non_tri:
        for a in s["starboard"] + s["port"]:
            if not a["dnf"] and a["split"] > 55:
                print(f"  WARN: {a['name']} split={a['split']} in {s['date']}")
                bad += 1
    if bad == 0:
        print("No splits > 55 in non-triathlon sessions")

    tri_count = sum(1 for s in sessions if s["workoutCategory"] == "triathlon")
    print(f"Triathlon sessions: {tri_count}")

    cats = Counter(s["workoutCategory"] for s in sessions)
    print(f"Categories: {dict(cats)}")

    print(f"\nTop 5:")
    for a in athletes[:5]:
        print(f"  #{a['rank']} {a['name']}: avg={a['avgSplit']}, best={a['bestSplit']}, n={a['sessionCount']}, trend={a['trend']}")

    # Generate TypeScript
    def ts(obj):
        return json.dumps(obj, indent=2)

    ts_lines = [
        "// AUTO-GENERATED by scripts/parse_erg_data.py — DO NOT EDIT",
        f"// {len(athletes)} athletes, {len(sessions)} sessions, Aug–Dec 2023",
        "",
        "export interface AthleteSplit {",
        "  name: string;",
        "  split: number;",
        "  spm: number;",
        "  rank: number;",
        "  dnf: boolean;",
        "  timeFormatted?: string;",
        "}",
        "",
        "export type WorkoutCategory = 'steady_state' | 'intervals' | 'threshold' | 'triathlon';",
        "",
        "export interface Session {",
        "  id: string;",
        "  date: string;",
        "  workoutName: string;",
        "  workoutCategory: WorkoutCategory;",
        "  starboard: AthleteSplit[];",
        "  port: AthleteSplit[];",
        "  teamAvgSplit: number;",
        "  bestSplit: number;",
        "  worstSplit: number;",
        "  bestAthlete: string;",
        "  starboardAvg: number;",
        "  portAvg: number;",
        "  participantCount: number;",
        "}",
        "",
        "export interface AthleteSession {",
        "  sessionId: string;",
        "  date: string;",
        "  workoutName: string;",
        "  workoutCategory: string;",
        "  side: 'starboard' | 'port';",
        "  split: number;",
        "  spm: number;",
        "  rank: number;",
        "  of: number;",
        "}",
        "",
        "export interface SparklinePoint { date: string; split: number; }",
        "",
        "export interface CategoryStats { sessions: number; avgSplit: number | null; }",
        "",
        "export interface Athlete {",
        "  id: string;",
        "  name: string;",
        "  group: 'S' | 'P';",
        "  rank: number;",
        "  avgSplit: number;",
        "  bestSplit: number;",
        "  bestSplitDate: string;",
        "  sessionCount: number;",
        "  trend: 'improving' | 'declining' | 'stable';",
        "  trendDelta: number;",
        "  consistency: number;",
        "  categoryBreakdown: Record<string, CategoryStats>;",
        "  sparklineData: SparklinePoint[];",
        "  sessions: AthleteSession[];",
        "}",
        "",
        f"export const sessions: Session[] = {ts(sessions)} as Session[];",
        "",
        f"export const athletes: Athlete[] = {ts(athletes)} as Athlete[];",
        "",
        f"export const sessionDates: string[] = {json.dumps([s['date'] for s in sessions])};",
        "",
    ]

    content = "\n".join(ts_lines)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(content)
    print(f"\nWritten {len(content)} bytes to {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
