#!/usr/bin/env python3
"""Parse Row IQ Dummy Erg Scores.xlsx into static JS data for the dashboard."""

import re
import json
import sys
from datetime import datetime, timedelta
from pathlib import Path
import openpyxl

EXCEL_PATH = Path(__file__).parent.parent.parent / "Downloads" / "Row IQ Dummy Erg Scores.xlsx"
OUTPUT_PATH = Path(__file__).parent.parent / "src" / "data" / "ergData.js"

# Name normalization map
NAME_FIXES = {
    'Pfautch': 'Pfautsch',
    'Andreucetti': 'Andreuccetti',
    'Distefano': 'DiStefano',
    'Unlod': 'Unold',
    'Joregensen': 'Jorgensen',
    'Mackinney': 'MacKinney',
    'Luludzija': 'Luledzija',
    'Ried': 'Reid',
    'Ferrrera': 'Ferrera',
    'Koudraeiv': 'Koudraiev',
    'Barrel': 'Barrell',
}

# Workout category classification
def classify_workout(workout_str):
    w = workout_str.lower().strip()
    if 'triathalon' in w or 'triathlon' in w:
        return 'triathlon'
    # Threshold: 5km, 10987, 5k4k3k2k1k, 7'30", 4x4', 302010
    if w in ['5km'] or ('5k' in w and '4k' in w):
        return 'threshold'
    if "10'" in w and "9'" in w and "8'" in w:  # 10987 pattern
        return 'threshold'
    if "30'/20'/10'" in w or '302010' in w:
        return 'threshold'
    if "7'" in w and '30' in w:
        return 'threshold'
    if "4x4'" in w and 'km' not in w:
        return 'threshold'
    # Intervals: 4x4km, 3x3km, 2x6km, 3x5.5km, 2x5km, 2x6k
    if re.search(r'\d+x\d+\.?\d*k', w):
        return 'intervals'
    # Steady state: 30', 2x20', 2x22', 3x15', 3x17', 6x10'
    if re.search(r"^\d+'", w) or re.search(r"\d+x\d+'", w) or re.search(r"\d+x\d+\'", w):
        return 'steady_state'
    if '30' in w and 'spm' in w:
        return 'steady_state'
    if '22' in w and 'rpm' in w:
        return 'steady_state'
    return 'steady_state'  # default

def normalize_name(name):
    if not name or not isinstance(name, str):
        return None
    n = name.strip()
    if not n:
        return None
    return NAME_FIXES.get(n, n)

SPLIT_RE = re.compile(r'(\d+\.?\d*)\s*\((\d+)\)\)*')

def parse_score(val):
    """Parse '44.7 (20)' -> (split=44.7, spm=20) or None."""
    if val is None:
        return None
    s = str(val).strip()
    if not s or s.lower() == 'nan':
        return None
    if 'dnf' in s.lower():
        return 'DNF'
    m = SPLIT_RE.search(s)
    if not m:
        return None
    split = float(m.group(1))
    spm = int(m.group(2))
    return (split, spm)

def find_sections(ws, header_row):
    """Scan a header row to find section boundaries."""
    sections = []
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=col).value
        if val is None:
            continue
        v = str(val).strip()
        vl = v.lower()
        if 'bike' in vl:
            sections.append(('bike', col, v))
        elif 'overall' in vl and 'avg' not in vl:
            sections.append(('overall', col, v))
        elif 'weight' in vl:
            sections.append(('weight_adjusted', col, v))
        elif vl == 'starboard':
            sections.append(('starboard', col, v))
        elif vl == 'port':
            sections.append(('port', col, v))
        elif 'erg' in vl:
            sections.append(('erg_sub', col, v))
        elif vl in ('5km', '5km '):
            # Separate 5km section in 5k4k3k2k1k sheet — this is erg data
            sections.append(('erg_extra', col, v))
    return sections

def find_avg_col(ws, header_row, section_start_col, next_section_col):
    """Find the AVG column within a section's range."""
    for col in range(section_start_col, min(next_section_col, ws.max_column + 1)):
        val = ws.cell(row=header_row, column=col).value
        if val is None:
            continue
        v = str(val).strip().lower()
        if v in ('avg', 'avg split', 'overall avg', 'time'):
            return col, v
    return None, None

def extract_athletes_from_section(ws, header_row, name_col, avg_col, score_type='split'):
    """Extract athlete data from a section."""
    athletes = []
    for row in range(header_row + 1, ws.max_row + 1):
        name_val = ws.cell(row=row, column=name_col).value
        name = normalize_name(name_val)
        if not name:
            continue
        # Skip sub-section headers that appear mid-data
        nl = name.lower()
        if nl in ('starboard', 'port', 'avg', 'bike', 'overall') or 'bike' in nl:
            break

        score_val = ws.cell(row=row, column=avg_col).value
        if score_type == 'time':
            # Triathlon time — skip
            continue

        result = parse_score(score_val)
        if result is None:
            continue
        if result == 'DNF':
            athletes.append({'name': name, 'split': None, 'spm': None, 'dnf': True})
            continue

        split, spm = result
        # Validation: discard bike data leakage
        if split > 60:
            continue
        athletes.append({'name': name, 'split': split, 'spm': spm, 'dnf': False})
    return athletes

def parse_session(ws):
    """Parse a single session worksheet."""
    # Get date and workout
    date_val = ws.cell(row=1, column=1).value
    workout_val = ws.cell(row=1, column=2).value

    if isinstance(date_val, datetime):
        date_str = date_val.strftime('%Y-%m-%d')
    else:
        return None

    workout = str(workout_val).strip() if workout_val else 'Unknown'
    category = classify_workout(workout)

    # Find header row — usually row 3, sometimes row 2
    header_row = 3
    sections_r3 = find_sections(ws, 3)
    sections_r2 = find_sections(ws, 2)

    if not any(s[0] in ('starboard', 'port') for s in sections_r3) and any(s[0] in ('starboard', 'port') for s in sections_r2):
        header_row = 2
        sections = sections_r2
    else:
        sections = sections_r3

    # If triathlon with no erg sub-section, check for "Erg" sub-section
    is_triathlon = 'triathalon' in workout.lower() or 'triathlon' in workout.lower()

    starboard_athletes = []
    port_athletes = []

    # Sort sections by column position
    sections.sort(key=lambda x: x[1])

    # Determine boundary columns for each section
    for i, (stype, scol, sname) in enumerate(sections):
        next_col = sections[i + 1][1] if i + 1 < len(sections) else ws.max_column + 1

        # Skip non-erg sections
        if stype in ('bike', 'overall', 'weight_adjusted'):
            continue

        if is_triathlon and stype in ('starboard', 'port'):
            # Triathlon starboard/port have Time (timedelta), not split — skip
            avg_col, avg_type = find_avg_col(ws, header_row, scol, next_col)
            if avg_type == 'time':
                continue

        if stype == 'erg_sub':
            # This is like "Erg 3x15'" in triathlon sheets — real erg data
            avg_col, _ = find_avg_col(ws, header_row, scol, next_col)
            if avg_col:
                athletes = extract_athletes_from_section(ws, header_row, scol, avg_col)
                # Add to starboard (arbitrary, just need them counted)
                starboard_athletes.extend(athletes)
            continue

        if stype == 'erg_extra':
            # Like "5km" extra section in 5k4k3k2k1k
            avg_col, _ = find_avg_col(ws, header_row, scol, next_col)
            if avg_col:
                athletes = extract_athletes_from_section(ws, header_row, scol, avg_col)
                port_athletes.extend(athletes)
            continue

        if stype in ('starboard', 'port'):
            # Find AVG column
            avg_col, avg_type = find_avg_col(ws, header_row, scol, next_col)
            if not avg_col:
                # Try col+1 as default
                avg_col = scol + 1

            if avg_type == 'time':
                continue  # Triathlon time data

            athletes = extract_athletes_from_section(ws, header_row, scol, avg_col)
            if stype == 'starboard':
                starboard_athletes.extend(athletes)
            else:
                port_athletes.extend(athletes)

    # Combine and rank
    all_athletes = []
    for a in starboard_athletes:
        a['group'] = 'Starboard'
        all_athletes.append(a)
    for a in port_athletes:
        a['group'] = 'Port'
        all_athletes.append(a)

    # Filter out DNF for ranking, but keep them in the list
    ranked = [a for a in all_athletes if not a['dnf'] and a['split'] is not None]
    ranked.sort(key=lambda x: x['split'])

    for i, a in enumerate(ranked):
        a['rank'] = i + 1

    for a in all_athletes:
        if a['dnf'] or a['split'] is None:
            a['rank'] = None

    # Compute summary stats
    splits = [a['split'] for a in ranked]
    if not splits:
        # Triathlon sessions with no erg data — still include the session
        if is_triathlon:
            return {
                'date': date_str,
                'workout': workout,
                'category': category,
                'starboard': [],
                'port': [],
                'teamAvg': None,
                'bestSplit': None,
                'bestAthlete': 'N/A',
                'participantCount': 0,
            }
        return None

    team_avg = round(sum(splits) / len(splits), 1)
    best_split = round(min(splits), 1)
    best_athlete = ranked[0]['name'] if ranked else 'N/A'

    # Separate back into groups for output
    sb_out = [{'name': a['name'], 'split': a['split'], 'spm': a['spm'], 'rank': a['rank'], 'dnf': a['dnf']}
              for a in all_athletes if a['group'] == 'Starboard']
    port_out = [{'name': a['name'], 'split': a['split'], 'spm': a['spm'], 'rank': a['rank'], 'dnf': a['dnf']}
                for a in all_athletes if a['group'] == 'Port']

    return {
        'date': date_str,
        'workout': workout,
        'category': category,
        'starboard': sb_out,
        'port': port_out,
        'teamAvg': team_avg,
        'bestSplit': best_split,
        'bestAthlete': best_athlete,
        'participantCount': len(all_athletes),
    }

def build_athlete_summaries(sessions):
    """Build per-athlete summary data across all sessions."""
    athlete_data = {}

    for sess in sessions:
        all_entries = sess['starboard'] + sess['port']
        total_in_session = len([e for e in all_entries if not e['dnf'] and e['split'] is not None])

        for entry in all_entries:
            name = entry['name']
            if name not in athlete_data:
                athlete_data[name] = {
                    'name': name,
                    'sessions': [],
                    'groups': [],
                }

            group = 'Starboard' if entry in sess['starboard'] else 'Port'
            athlete_data[name]['groups'].append(group)
            athlete_data[name]['sessions'].append({
                'date': sess['date'],
                'workout': sess['workout'],
                'category': sess['category'],
                'split': entry['split'],
                'spm': entry['spm'],
                'rank': entry['rank'],
                'of': total_in_session,
                'dnf': entry['dnf'],
            })

    # Compute summaries
    summaries = []
    for name, data in athlete_data.items():
        valid_sessions = [s for s in data['sessions'] if not s['dnf'] and s['split'] is not None]
        if not valid_sessions:
            continue

        splits = [s['split'] for s in valid_sessions]
        avg_split = round(sum(splits) / len(splits), 1)
        best_split = round(min(splits), 1)

        # Primary group = most common
        from collections import Counter
        group_counts = Counter(data['groups'])
        primary_group = group_counts.most_common(1)[0][0]

        # Trend: first 3 vs last 3 sessions (need ≥6)
        trend = 'stable'
        trend_delta = 0
        if len(valid_sessions) >= 6:
            first3 = sum(s['split'] for s in valid_sessions[:3]) / 3
            last3 = sum(s['split'] for s in valid_sessions[-3:]) / 3
            trend_delta = round(last3 - first3, 1)
            if trend_delta < -0.5:
                trend = 'improving'
            elif trend_delta > 0.5:
                trend = 'declining'

        summaries.append({
            'name': name,
            'avgSplit': avg_split,
            'bestSplit': best_split,
            'totalSessions': len(data['sessions']),
            'trend': trend,
            'trendDelta': trend_delta,
            'primaryGroup': primary_group,
            'history': data['sessions'],
        })

    # Rank by avgSplit
    summaries.sort(key=lambda x: x['avgSplit'])
    for i, s in enumerate(summaries):
        s['rank'] = i + 1

    return summaries

def main():
    print(f"Loading {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=True)
    print(f"Found {len(wb.sheetnames)} sheets: {wb.sheetnames}")

    sessions = []
    for i, sheet_name in enumerate(wb.sheetnames):
        if i == 0:  # Skip "Erg Log"
            continue
        ws = wb[sheet_name]
        print(f"\nParsing sheet {i}: '{sheet_name}'...")
        result = parse_session(ws)
        if result:
            sessions.append(result)
            print(f"  -> {result['date']} | {result['workout']} | {result['category']} | "
                  f"{result['participantCount']} athletes | avg={result['teamAvg']} | "
                  f"best={result['bestSplit']} ({result['bestAthlete']})")
        else:
            print(f"  -> SKIPPED (no valid erg data)")

    print(f"\n{'='*60}")
    print(f"Total sessions parsed: {len(sessions)}")

    # Build athlete summaries
    athletes = build_athlete_summaries(sessions)
    print(f"Total unique athletes: {len(athletes)}")

    # Validation
    print(f"\n--- VALIDATION ---")
    print(f"Sessions: {len(sessions)} (expected 28)")
    print(f"Date range: {sessions[0]['date']} to {sessions[-1]['date']}")
    print(f"Athletes: {len(athletes)}")
    print(f"\nTop 10 athletes:")
    for a in athletes[:10]:
        print(f"  #{a['rank']} {a['name']}: avg={a['avgSplit']}s, best={a['bestSplit']}s, "
              f"sessions={a['totalSessions']}, trend={a['trend']} ({a['trendDelta']:+.1f})")

    # Check for any avg > 55
    bad = [a for a in athletes if a['avgSplit'] > 55]
    if bad:
        print(f"\n⚠️  Athletes with avg > 55s (likely data errors):")
        for a in bad:
            print(f"  {a['name']}: {a['avgSplit']}s")
    else:
        print(f"\n✅ No athletes with avg > 55s")

    # Check Phelps ranking
    phelps = next((a for a in athletes if a['name'] == 'Phelps'), None)
    if phelps:
        print(f"\n{'✅' if phelps['rank'] <= 2 else '⚠️'} Phelps rank: #{phelps['rank']} (avg={phelps['avgSplit']}s)")
    else:
        print(f"\n⚠️  Phelps NOT FOUND!")

    # Write output
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    js_content = "// Auto-generated from Row IQ Dummy Erg Scores.xlsx\n"
    js_content += "// Do not edit manually — run scripts/parse_erg_data.py to regenerate\n\n"
    js_content += f"export const sessions = {json.dumps(sessions, indent=2)};\n\n"
    js_content += f"export const athletes = {json.dumps(athletes, indent=2)};\n\n"

    # Category summary
    categories = {}
    for s in sessions:
        cat = s['category']
        if cat not in categories:
            categories[cat] = {'name': cat, 'count': 0, 'totalAvg': 0, 'avgCount': 0, 'bestSplit': 999, 'bestAthlete': ''}
        categories[cat]['count'] += 1
        if s['teamAvg'] is not None:
            categories[cat]['totalAvg'] += s['teamAvg']
            categories[cat]['avgCount'] = categories.get(cat, {}).get('avgCount', 0) + 1
        if s['bestSplit'] is not None and s['bestSplit'] < categories[cat]['bestSplit']:
            categories[cat]['bestSplit'] = s['bestSplit']
            categories[cat]['bestAthlete'] = s['bestAthlete']

    for cat in categories.values():
        cat['teamAvg'] = round(cat['totalAvg'] / max(cat['avgCount'], 1), 1)
        del cat['totalAvg']
        del cat['avgCount']

    js_content += f"export const categoryStats = {json.dumps(list(categories.values()), indent=2)};\n"

    with open(OUTPUT_PATH, 'w') as f:
        f.write(js_content)

    print(f"\n✅ Written to {OUTPUT_PATH}")
    print(f"   {len(sessions)} sessions, {len(athletes)} athletes, {len(categories)} categories")

if __name__ == '__main__':
    main()
